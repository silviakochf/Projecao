"""Excel da progressão de alunos (a partir do resultado da análise automática)."""
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AZUL, AZUL2, CINZA = '0A192F', '172A45', '334155'
VERDE, AMARELO, LARANJA, VERMELHO, ROXO, AZUL_CL = 'D1FAE5', 'FEF9C3', 'FFEDD5', 'FEE2E2', 'EDE9FE', 'DBEAFE'
COR_PROX = {'mesmo': VERDE, 'proximo': AMARELO, 'distante': LARANJA}


def _fill(h): return PatternFill('solid', fgColor=h)
def _font(b=False, c='1E293B', s=9): return Font(name='Arial', bold=b, color=c, size=s)
def _borda():
    l = Side(style='thin', color='CBD5E1')
    return Border(left=l, right=l, top=l, bottom=l)


def _titulo(ws, ate, texto, sub):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(f'A1:{ate}1'); ws.merge_cells(f'A2:{ate}2')
    ws['A1'].value, ws['A1'].font, ws['A1'].fill = texto, _font(True, 'FFFFFF', 13), _fill(AZUL)
    ws['A2'].value, ws['A2'].font, ws['A2'].fill = sub, _font(False, 'FFFFFF', 9), _fill(AZUL2)
    for c in ('A1', 'A2'):
        ws[c].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 18


def _tabela(ws, linha, cabecalho, linhas, larguras, cor_linha=None, centro=()):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, t in enumerate(cabecalho, 1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font, c.fill, c.border = _font(True, 'FFFFFF'), _fill(CINZA), _borda()
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[linha].height = 28
    for n, valores in enumerate(linhas, linha + 1):
        cor = cor_linha(valores) if cor_linha else None
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=n, column=i, value=v)
            c.font, c.border = _font(), _borda()
            if cor:
                c.fill = _fill(cor)
            c.alignment = Alignment(horizontal='center' if (i - 1) in centro else 'left', vertical='center')
        ws.row_dimensions[n].height = 16
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)
    ws.auto_filter.ref = f'A{linha}:{get_column_letter(len(cabecalho))}{linha + max(len(linhas), 1)}'


def gerar_excel_progressao(resultado: dict) -> io.BytesIO:
    turmas, destinos, r = resultado['turmas'], resultado.get('destinos', []), resultado['resumo']
    agora = datetime.now().strftime('%d/%m/%Y %H:%M')
    sub = (f"Gerado em {agora}  |  {r['turmas']} turmas / {r['alunos']} alunos  |  "
           f"alocadas: {r['turmas_alocadas']} ({r['estadual']} estadual, {r['municipal']} municipal)  |  "
           f"sem vaga: {r['sem_vaga']} turmas / {r['alunos_sem_vaga']} alunos")
    wb = Workbook()

    # 1) Progressão — uma linha por turma
    ws = wb.active; ws.title = 'Progressão'
    linhas, cores = [], []
    for n, t in enumerate(turmas, 1):
        s = t.get('escola_sugerida')
        linhas.append([n, t['escola_origem'], t['bairro_origem'] or 'não identificado', t['turma'],
                       t['serie_atual'], t['serie_destino'], t['turno'], t['total_alunos'],
                       s['escola'] if s else '— SEM VAGA PARA A TURMA INTEIRA —',
                       s['tipo'] if s else '', s['bairro'] if s else '',
                       s['proximidade_rotulo'] if s else '', s['vagas_disponiveis'] if s else '',
                       ('turno diferente do atual' if s and s['turno_diferente'] else '') if s else 'ver aba "Sem vaga"'])
    _titulo(ws, 'N', 'PROGRESSÃO DE ALUNOS — SUGESTÃO AUTOMÁTICA', sub)
    def cor(v):
        if not v[9]: return VERMELHO
        return COR_PROX.get({'Mesmo bairro': 'mesmo', 'Bairro próximo': 'proximo'}.get(v[11], 'distante'))
    _tabela(ws, 4, ['#', 'Escola de origem', 'Bairro origem', 'Turma', 'Série atual', 'Série destino', 'Turno',
                    'Alunos', 'Escola de destino', 'Tipo', 'Bairro destino', 'Proximidade',
                    'Vagas do destino (antes)', 'Observação'],
            linhas, [5, 42, 18, 14, 10, 12, 10, 8, 44, 12, 18, 16, 14, 24], cor, centro=(0, 4, 5, 6, 7, 9, 11, 12))

    # 2) Por escola de destino
    wd = wb.create_sheet('Por escola de destino')
    grupos = {}
    for t in turmas:
        s = t.get('escola_sugerida')
        if not s: continue
        g = grupos.setdefault((s['tipo'] != 'ESTADUAL', s['escola'], t['serie_destino'], s['turno']),
                              {'bairro': s['bairro'], 'tipo': s['tipo'], 'turmas': 0, 'alunos': 0, 'orig': []})
        g['turmas'] += 1; g['alunos'] += t['total_alunos']
        if t['escola_origem'] not in g['orig']: g['orig'].append(t['escola_origem'])
    linhas = [[g['tipo'], k[1], g['bairro'], k[2], k[3], g['turmas'], g['alunos'], '; '.join(g['orig'])]
              for k, g in sorted(grupos.items())]
    _titulo(wd, 'H', 'ALOCAÇÕES POR ESCOLA DE DESTINO', sub)
    _tabela(wd, 4, ['Tipo', 'Escola de destino', 'Bairro', 'Série', 'Turno', 'Turmas', 'Alunos', 'Vêm de'],
            linhas, [12, 44, 18, 10, 11, 8, 8, 80], lambda v: ROXO if v[0] == 'ESTADUAL' else AZUL_CL,
            centro=(0, 3, 4, 5, 6))

    # 3) Sem vaga
    wv = wb.create_sheet('Sem vaga')
    linhas = []
    for t in turmas:
        if t.get('escola_sugerida'): continue
        alt = '; '.join(f"{c['escola']} ({c['vagas_disponiveis']} vagas)" for c in t['sugestoes']) \
              or 'nenhuma unidade oferece essa série'
        linhas.append([t['escola_origem'], t['bairro_origem'] or 'não identificado', t['turma'],
                       t['serie_atual'], t['serie_destino'], t['turno'], t['total_alunos'], alt])
    _titulo(wv, 'H', 'TURMAS SEM VAGA PARA A TURMA INTEIRA', sub)
    _tabela(wv, 4, ['Escola de origem', 'Bairro', 'Turma', 'Série atual', 'Série destino', 'Turno',
                    'Alunos', 'Maiores vagas encontradas (insuficientes p/ a turma inteira)'],
            linhas, [42, 18, 14, 10, 12, 10, 8, 100], lambda v: VERMELHO, centro=(3, 4, 5, 6))

    # 4) Projeção de vagas
    wp = wb.create_sheet('Vagas (projeção)')
    linhas = []
    for c in sorted(destinos, key=lambda c: (c['tipo'], c['escola'], c['serie'], c['turno'])):
        linhas.append([c['escola'], c['tipo'], c['bairro'] or 'não identificado', c['serie'], c['turno'],
                       c['total'] if c['total'] is not None else '—',
                       c['ocupadas'] if c['ocupadas'] is not None else '—',
                       c['livres_hoje'] if c['livres_hoje'] is not None else '—',
                       c['fluxo'], c['disponiveis'], c['alocados'], c['restantes']])
    _titulo(wp, 'L', 'PROJEÇÃO DE VAGAS PARA O ANO SEGUINTE', sub)
    _tabela(wp, 4, ['Escola', 'Tipo', 'Bairro', 'Série', 'Turno', 'Capacidade total (ocup.+livres)',
                    'Ocupadas hoje', 'Livres hoje', 'Sobem da série anterior (mesma escola)',
                    'Vagas ano seguinte', 'Alocadas pela IA', 'Restantes'],
            linhas, [44, 12, 18, 10, 11, 16, 11, 10, 18, 12, 12, 11],
            lambda v: ROXO if v[1] == 'ESTADUAL' else None, centro=(1, 3, 4, 5, 6, 7, 8, 9, 10, 11))

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf
